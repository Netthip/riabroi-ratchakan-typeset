# -*- coding: utf-8 -*-
"""
ตัดคำไทยสำหรับเอกสารราชการ — v1.1

ทำงาน 2 ชั้น เหมือนเครื่องกำเนิดใบเสนองาน (โม 09) ของกิ๊ฟ
  ชั้น 1  หาจุดตัดระดับคำ แล้วใส่ ZWSP (U+200B) ตรงขอบเขตคำ
  ชั้น 2  keep_together — ลบจุดตัดที่ทำให้อ่านสะดุด และมัดคำด้วย NBSP (U+00A0)

ชั้น 1 ต่างจาก thai-docx (ตัวสาธารณะ) ตรงที่
  - ตัดคำจาก "ข้อความทั้งย่อหน้า" แล้วค่อยแมปกลับเข้า run
    เอกสารจาก Word มักมี run ที่ผ่ากลางคำ (เจอจริง: 'ประชาชนไม่โด' + 'นทอดทิ้ง')
  - NO_BREAK_BEFORE ไม่รวมสระหน้า เ แ โ ใ ไ — ถ้ารวม คำอย่าง และ/ได้/เป็น จะเกาะคำก่อนหน้า
    เป็นก้อนยาว 25-30 ตัวจนบรรทัดถ่าง
  - min_token = 2 · ซอยคำประสมยาวเกิน 12 ตัวตรงจุดที่ได้คำจริงทั้งสองฝั่ง
  - ใส่ ZWSP เฉพาะจุด "ไทยชนไทย" ไม่แหกคำอังกฤษ ตัวเลข เบอร์โทร

ชั้น 2 ถอดจาก thai_text.keep_together ของโม 09 (ก.ย. 2569) แล้วเพิ่มกฎจากการอ่านไล่
รอยต่อบรรทัดทีละจุดตอนทำตัวอย่างเอกสารประเด็นถามตอบ (21 ก.ย. 2569) — ดู _reading_rules

ใส่ ZWSP เฉพาะเนื้อความที่จัดชิดขอบขวาเท่านั้น (insert_zwsp_document ทำให้อัตโนมัติ)
บทเรียนโม 09 (11 ก.ค. 2569): เคยแทรก ZWSP ทุกคำทุกย่อหน้า กิ๊ฟบ่นว่า "แยกถูกผิดไม่ออก"

หมายเหตุสำหรับคนแก้ไฟล์นี้: อย่าพิมพ์อักขระ ZWSP/NBSP ลงซอร์สตรง ๆ และอย่าเขียน escape ยูนิโค้ด
ผ่านเครื่องมือแก้ไฟล์ เพราะจะกลายเป็นตัวล่องหนที่มองไม่เห็นใน diff — ใช้ตัวแปร ZWSP/NBSP เสมอ
"""
from __future__ import annotations

import re
import unicodedata
import warnings

ZWSP = chr(0x200B)
NBSP = chr(0x00A0)

# ─────────────────────────────────────────────────────────── ตัวอักษรไทย

THAI_CHAR = re.compile("[" + chr(0x0E00) + "-" + chr(0x0E7F) + "]")
_TH = "[" + chr(0x0E01) + "-" + chr(0x0E4E) + "]"          # พยัญชนะถึงเครื่องหมาย (ไม่รวมเลขไทย)
_DIG = "[0-9๐-๙]"

# อักขระที่ห้ามขึ้นต้นบรรทัด = สระบน/ล่าง/หลัง วรรณยุกต์ ไม้ยมก ฯลฯ
# เจตนา: ไม่มีสระหน้า (เ แ โ ใ ไ) อยู่ในชุดนี้
NO_BREAK_BEFORE = set("ๆฯะัาำิีึืุูๅ็่้๊๋์ํฺ")

# วลีที่ห้าม ZWSP ไปตัดกลาง — ตัวตัดคำมองเป็นคนละคำ แต่ในหนังสือราชการต้องอยู่ด้วยกัน
# ใส่เฉพาะที่กิ๊ฟสั่งจริง ห้ามเติมวลียาวเอง: ยิ่งมัดยาว จุดตัดยิ่งน้อย Word จะยืดอักษรจนผิดปกติ
# (เคยใส่ 'แผนการปฏิบัติงาน' แล้วช่องไฟพุ่งเป็น 7 pt)
KEEP_TOGETHER = (
    "ประจำปีงบประมาณ", "ตามที่ขอทำความตกลงไปได้",
    "ความละเอียด", "สงป.",
    "แผนงาน", "โครงการ",
    "ตัวชี้วัด", "ก่อหนี้", "แนวทาง", "รวมทั้งสิ้น",
)

# คำจริงที่พจนานุกรม pythainlp ไม่มี ตัวตัดคำจึงผ่ากลาง — เจอจากการอ่านไล่รอยต่อบรรทัด
# ต่างจาก KEEP_TOGETHER ตรงที่เป็น "คำ" ไม่ใช่วลี เจอคำใหม่แบบนี้ เพิ่มได้เลย
DICT_GAPS = ("เฝ้าระวัง", "กรมควบคุมโรค", "รวมถึง", "สัญชาติไทย",
             "น้ำหนักสัมพัทธ์", "เทียบเท่าเงินสด")

# คำประสมยาวเกิน MAX_CHUNK ที่ห้ามซอย เพราะซอยแล้วแต่ละครึ่งอ่านเป็นคนละเรื่อง
# '(โทรศัพท์ / มือถือ/แท็บเล็ต)' อ่านเป็นรายการ 3 อย่าง · 'รวมถึงสิทธิ / ประโยชน์อื่น ๆ'
# ต่างจาก 'งบประมาณ|รายจ่าย' ที่ซอยแล้วยังอ่านเป็นเรื่องเดียวกัน (เจอตอนทดสอบ v1.1 21 ก.ย. 2569)
WHOLE_COMPOUNDS = ("โทรศัพท์มือถือ", "สิทธิประโยชน์", "จัดซื้อจัดจ้าง")

_KEEP_WHOLE = KEEP_TOGETHER + DICT_GAPS + WHOLE_COMPOUNDS

# คำนำวลีที่ห้ามค้างท้ายบรรทัด ต้องลงไปขึ้นบรรทัดใหม่พร้อมคำที่ตามมา
# หลักของกิ๊ฟไม่ใช่รายการคำบังคับ แต่คือ "คำที่ค้างบรรทัดบนเดี่ยว ๆ แล้วอ่านสะดุด"
LEAD_WORDS = ("เรื่อง", "ให้", "เมื่อ", "ที่เป็น", "ตาม", "มี", "ต่อ",
              "วงเงิน", "แผนการ", "สถาบัน", "เพื่อ")
# คำนำที่เป็นคำนาม เป็นกรรมของคำนำตัวหน้าได้ จึงติดกันเป็นทอด (มี/วงเงิน · ให้/สถาบัน)
# คำนำที่เหลือเป็นคำเชื่อม/บุพบท = ต้นวลีใหม่ ตัดข้างหน้าได้ (…ให้เต็มวงเงิน / เมื่อสำนักงบประมาณ)
LEAD_NOUNS = ("วงเงิน", "แผนการ", "สถาบัน")
# ติดคำถัดไปเฉพาะเมื่อตามด้วย LEAD_NOUNS — '…๑ โครงการ รวม / วงเงิน ๕๓,๓๙๓,๔๐๐ บาท'
LEAD_BEFORE_NOUNS = ("รวม",)
# ห้ามขึ้นต้นบรรทัด ต้องติดคำข้างหน้า — 'สถาบันวัคซีน / แห่งชาติ'
TAIL_WORDS = ("แห่งชาติ",)
# คำที่กิ๊ฟเองปล่อยให้ค้างท้ายบรรทัดได้ ไม่ต้องมัด (ไว้อ้างอิง — ไม่ได้อยู่ใน LEAD_WORDS)
ALLOW_DANGLING = ("และ", "หรือ", "ของ", "ใน", "ตั้งแต่")

# คำนำที่ประสมกันเป็นคำเดียว ตัวตัดคำจึงไม่เห็นขอบคำหน้าคำนำตัวหลัง กฎคำนำปกติเลยไม่ทำงาน
# 'โรคให้ / มีประสิทธิภาพ' · 'ของการมี / สมุดบันทึก' · 'อย่าง / มีประสิทธิภาพ'
COMPOUND_LEADS = ("ให้มี", "การมี", "อย่างมี", "อย่างเป็น")
# คำนำหน้าที่ยืนเดี่ยวแล้วต้องติดคำถัดไป — 'สนับสนุนเชิง / นโยบาย' · 'ติดต่อทาง / พันธุกรรม'
PREFIX_WORDS = ("เชิง", "ทาง")
# หน่วยนับที่ต้องลงมาพร้อมตัวเลข — "จำนวน ๔ แผนงาน"
UNITS = ("แผนงาน", "ผลผลิต", "โครงการ", "รายการ")

MAX_CHUNK = 12          # ก้อนยาวกว่านี้ให้หาจุดตัดย่อยเพิ่ม
MIN_TOKEN = 2           # ชิ้นสั้นกว่านี้เกาะคำก่อนหน้า (2 สำหรับงานจัดชิดขอบขวา)

# คำที่มักถูกตัวตัดคำผลักไปเกาะคำถัดไปทั้งที่เป็นหางของคำหน้า
# 'ดำเนิน|การตาม' ต้องเป็น 'ดำเนินการ|ตาม' — เช็กกับพจนานุกรมก่อนย้ายเสมอ
NOMINAL_PREFIX = ("การ", "ความ")

_TOKENIZER = None
_WORDS = None

# ขอบคำ = ต้นข้อความ · ช่องว่าง · ZWSP · NBSP · วงเล็บเปิด (lookbehind ที่ไม่กินตัวอักษร)
_B = r"(?<![^\s" + ZWSP + NBSP + r"(])"
_SEP = "[ " + ZWSP + "]+"          # ช่องว่างหรือ ZWSP อย่างน้อยหนึ่งตัว
_Z = ZWSP + "+"


def _tokenize(text: str) -> list[str]:
    """
    ตัดคำด้วย pythainlp ถ้าไม่มีจะคืนทั้งก้อน (ไม่พัง แต่ได้จุดตัด 0 จุด)

    กรณีนั้นต้องเตือนดัง ๆ — sandbox ของผู้ช่วย AI บางตัวไม่มี pythainlp และติดตั้งเพิ่มไม่ได้
    ถ้าเงียบ ผู้ใช้จะได้ไฟล์ที่ยังตัดกลางคำอยู่ โดยที่ทั้งคนและผู้ช่วยคิดว่าแก้แล้ว
    """
    global _TOKENIZER
    if _TOKENIZER is None:
        try:
            from pythainlp.tokenize import word_tokenize
            _TOKENIZER = word_tokenize
        except ImportError:
            _TOKENIZER = False
            warnings.warn(
                "ไม่พบ pythainlp — ตัดคำไทยไม่ได้ จะไม่มีการใส่ ZWSP เลย "
                "(ติดตั้ง: pip install pythainlp) | pythainlp not installed: "
                "Thai word breaking is disabled and no ZWSP will be inserted",
                RuntimeWarning, stacklevel=2)
    if not _TOKENIZER:
        return [text]
    return _TOKENIZER(text, engine="newmm")


def _words() -> set:
    """พจนานุกรมไว้ยืนยันว่าจุดที่จะซอย/ย้าย ได้คำจริงทั้งสองฝั่ง"""
    global _WORDS
    if _WORDS is None:
        try:
            from pythainlp.corpus import thai_words
            _WORDS = set(thai_words())
        except ImportError:
            _WORDS = set()
    return _WORDS


def is_thai(ch: str) -> bool:
    return bool(THAI_CHAR.match(ch))


# ═══════════════════════════════════════════ ชั้น 1: หาจุดตัดระดับคำ

def _best_split(tok: str) -> int:
    """
    หาจุดซอยที่ดีที่สุดในคำยาว: เอาจุดที่ 'ทั้งสองฝั่งเป็นคำจริง' และใกล้กลางที่สุด
    'งบประมาณรายจ่าย' -> งบประมาณ|รายจ่าย ไม่ใช่ งบประมาณราย|จ่าย
    ถ้าไม่มีจุดไหนได้คำจริงทั้งคู่ คืน 0 (แปลว่าอย่าซอย ปล่อยยาวดีกว่าซอยมั่ว)
    """
    words = _words()
    n = len(tok)
    mid = n // 2
    for off in range(0, n):
        for cut in {mid - off, mid + off}:
            if not (2 <= cut <= n - 2):
                continue
            if tok[cut] in NO_BREAK_BEFORE or unicodedata.combining(tok[cut]):
                continue
            if tok[:cut] in words and tok[cut:] in words:
                return cut
    return 0


def split_long_chunks(tokens: list[str], max_chars: int = MAX_CHUNK) -> list[str]:
    """
    ซอยคำประสมยาว ๆ ที่พจนานุกรมมองเป็นคำเดียว — 'งบประมาณรายจ่าย' -> 'งบประมาณ' + 'รายจ่าย'
    ซอยเฉพาะจุดที่ได้คำจริงทั้งสองฝั่ง และไม่แตะวลี/คำที่ตั้งใจมัดไว้
    """
    out: list[str] = []
    for tok in tokens:
        if any(p in tok for p in _KEEP_WHOLE):
            out.append(tok)
            continue
        while len(tok) > max_chars:
            cut = _best_split(tok)
            if not cut:
                break
            out.append(tok[:cut])
            tok = tok[cut:]
        out.append(tok)
    return out


def _fix_prefix_split(tokens: list[str]) -> list[str]:
    """
    ย้าย 'การ'/'ความ' ที่ตัวตัดคำผลักไปเกาะคำถัดไป กลับมาเป็นหางของคำหน้า
    'ดำเนิน' + 'การตาม'  ->  'ดำเนินการ' + 'ตาม'   (เพราะ 'ดำเนินการ' เป็นคำจริง)
    """
    words = _words()
    out = list(tokens)
    for i in range(len(out) - 1):
        nxt = out[i + 1]
        for pre in NOMINAL_PREFIX:
            if nxt.startswith(pre) and len(nxt) > len(pre):
                if (out[i] + pre) in words:
                    out[i] = out[i] + pre
                    out[i + 1] = nxt[len(pre):]
                break
    return [t for t in out if t]


def _merge_small(tokens: list[str], min_token: int) -> list[str]:
    """ชิ้นที่สั้นกว่า min_token ให้เกาะคำก่อนหน้า กันตัดกลางคำประสม"""
    out: list[str] = []
    for tok in tokens:
        if out and len(tok.strip()) < min_token and is_thai(tok[:1] or " "):
            out[-1] += tok
        else:
            out.append(tok)
    return out


def _protected_spans(text: str) -> list[tuple[int, int]]:
    """
    ช่วงตัวอักษรที่ห้ามมีจุดตัดอยู่ข้างใน — ทำที่ระดับตัวอักษร ไม่ใช่ระดับ token
    เพราะขอบเขต token ไม่ตรงกับวลีเสมอ (ตัวตัดคำคืน 'ตามตัว|ชี้|วัด' วลีจึงคร่อมกลาง token)
    """
    spans = []
    for phrase in _KEEP_WHOLE:
        start = 0
        while True:
            k = text.find(phrase, start)
            if k < 0:
                break
            spans.append((k, k + len(phrase)))
            start = k + 1
    return spans


def break_offsets(text: str, *, min_token: int = MIN_TOKEN,
                  max_chunk: int = MAX_CHUNK) -> set[int]:
    """ตำแหน่ง (ดัชนีตัวอักษร) ที่ใส่ ZWSP ได้ — ชั้น 1 เท่านั้น ยังไม่ผ่าน keep_together"""
    if not text or not THAI_CHAR.search(text):
        return set()

    toks = _tokenize(text)
    toks = _fix_prefix_split(toks)
    toks = split_long_chunks(toks, max_chunk)
    toks = _merge_small(toks, min_token)

    cuts, pos = set(), 0
    for tok in toks:
        pos += len(tok)
        cuts.add(pos)
    cuts.discard(0)
    cuts.discard(len(text))

    spans = _protected_spans(text)
    ok = set()
    for c in cuts:
        prev, nxt = text[c - 1], text[c]
        if not (is_thai(prev) and is_thai(nxt)):
            continue            # ไทยชนไทยเท่านั้น จะได้ไม่แหกคำอังกฤษ/ตัวเลข
        if nxt in NO_BREAK_BEFORE or unicodedata.combining(nxt):
            continue            # ห้ามขึ้นบรรทัดด้วยสระบน/ล่าง/วรรณยุกต์
        if any(s < c < e for s, e in spans):
            continue            # อยู่กลางวลี/คำที่ห้ามขาด
        ok.add(c)
    return ok


def _insert_raw(text: str, **kw) -> str:
    cuts = break_offsets(text, **kw)
    if not cuts:
        return text
    out, prev = [], 0
    for c in sorted(cuts):
        out.append(text[prev:c])
        prev = c
    out.append(text[prev:])
    return ZWSP.join(out)


# ═══════════════════════════════════ ชั้น 2: keep_together (จากโม 09)

def _lead(word: str) -> str:
    # ตัวตัดคำแยก "ที่เป็น" เป็น "ที่|เป็น" จึงยอม ZWSP คั่นกลางได้
    return "ที่" + ZWSP + "?เป็น" if word == "ที่เป็น" else re.escape(word)


def _glue_lead_words(text: str) -> str:
    """
    ลบจุดตัดบรรทัด 'หลัง' คำนำวลี และ 'ก่อน' คำท้าย — ช่องว่างเปลี่ยนเป็น NBSP จำนวนเท่าเดิม

    ติดเฉพาะเมื่อคำนำยืนเป็นคำเดี่ยว ("ติดตาม" ไม่โดนเพราะ "ตาม" ติดอยู่กับ "ติด")
    และไม่ติดถ้าคำถัดไปเป็นคำนำชนิดคำเชื่อม ซึ่งคือต้นวลีใหม่
    (โม 09 เคยติดหมดจน "ให้เต็มวงเงิน / เมื่อสำนักงบประมาณ" กลายเป็นก้อนเดียวตัดไม่ได้)
    """
    def glue(m):
        return m.group(1).replace(ZWSP, "") + m.group(2).replace(ZWSP, "").replace(" ", NBSP)

    heads = "|".join(_lead(w) for w in sorted(LEAD_WORDS, key=len, reverse=True))
    starts = "|".join(_lead(w) for w in LEAD_WORDS if w not in LEAD_NOUNS)
    text = re.sub(_B + "(" + heads + ")(" + _SEP + r")(?=\S)(?!(?:" + starts + "))", glue, text)
    # "รวม / วงเงิน…" ทำหลังรอบบนเสมอ
    nouns = "|".join(_lead(w) for w in LEAD_NOUNS)
    text = re.sub(_B + "(" + "|".join(map(re.escape, LEAD_BEFORE_NOUNS)) + ")(" + _SEP
                  + ")(?=(?:" + nouns + "))", glue, text)
    # "วงเงิน(งบประมาณ)" ต้องอยู่บรรทัดเดียวกับจำนวนเงินที่ตามมา — ไม่ครอบ "…รวมทั้งสิ้น"
    # เพราะก้อนยาวเกือบครึ่งบรรทัด ตกลงบรรทัดล่างทีไรบรรทัดบนถูกถ่างทั้งแถว
    text = re.sub("(วงเงิน(?:" + ZWSP + "?งบประมาณ)?)" + _SEP + "(" + _DIG + "[0-9๐-๙,.]*" + NBSP + "บาท)",
                  lambda m: m.group(1).replace(ZWSP, "") + NBSP + m.group(2), text)
    for word in TAIL_WORDS:
        text = re.sub(_Z + "(" + re.escape(word) + ")", r"\1", text)
    return text


def _phrase_pattern(phrase: str) -> str:
    """ZWSP ได้เฉพาะ 'ระหว่าง' ตัวอักษรของวลี — ไม่กิน ZWSP ที่อยู่หลังวลี
    (โม 09 ต่อ ZWSP? หลังตัวสุดท้ายด้วย จุดตัดหลังวลีจึงหายไปทุกครั้งโดยไม่ได้ตั้งใจ)"""
    return (ZWSP + "?").join(re.escape(ch) for ch in phrase)


def _reading_rules(text: str) -> str:
    """
    กฎเพิ่มจากการอ่านไล่รอยต่อบรรทัดทีละจุด (ตัวอย่างประเด็นถามตอบ 21 ก.ย. 2569)
    ทุกข้อเป็นแบบ "ห้ามตัดตรงนี้" — ไม่เปลี่ยนข้อความ
    """
    # ชื่อย่อ ≤6 ตัวที่ลงท้ายด้วยจุด ห้ามตัดข้างใน (สว|รส. -> สวรส.)
    text = re.sub(_B + "((?:" + _TH + ZWSP + "?){1,6})(?=\\.)",
                  lambda m: m.group(1).replace(ZWSP, ""), text)
    # จุดที่ต่อด้วยอักษรไทยทันที = ชื่อย่อ ห้ามตัดหลังจุด (ม.|มหิดล) — จบประโยคจะมีช่องว่างตามเสมอ
    text = re.sub("\\." + _Z + "(?=" + _TH + ")", ".", text)
    # คำนำประสม: ให้มี · การมี · อย่างมี · อย่างเป็น · เพื่อให้ -> ติดคำถัดไป
    text = re.sub("ให้" + _Z + "(?=มี)", "ให้", text)
    text = re.sub("อย่าง" + _Z + "(?=มี|เป็น)", "อย่าง", text)
    for head in COMPOUND_LEADS:
        text = re.sub("(?<=" + head + ")" + _Z, "", text)
    starts = "|".join(_lead(w) for w in LEAD_WORDS if w not in LEAD_NOUNS)
    text = re.sub("(?<=เพื่อให้)" + _Z + "(?!(?:" + starts + "))", "", text)
    # คำนำหน้าที่ยืนเดี่ยว ติดคำถัดไป (ต้องอยู่ที่ขอบคำ — "ทาง" ท้าย "แนวทาง" ไม่โดน)
    for w in PREFIX_WORDS:
        text = re.sub(_B + w + _Z, w, text)
    return text


def keep_together(text: str) -> str:
    """
    ชั้น 2 — ลบจุดตัดที่อ่านสะดุด และมัดคำด้วย NBSP
    ถอดจาก Gift Command Center โม 09 thai_text.keep_together แล้วเพิ่ม _reading_rules

    เปลี่ยนได้แค่ 2 อย่าง: ลบ ZWSP และเปลี่ยนช่องว่างเป็น NBSP — ตัวอักษรอื่นต้องเหมือนเดิม
    """
    if not text:
        return text or ""
    for phrase in _KEEP_WHOLE:
        text = re.sub(_phrase_pattern(phrase), phrase, text)
    # หน่วยนับลงมาทั้งก้อน — "จำนวน ๔ แผนงาน" · "รวม ๔ แผนงาน" (เลขไม่เกิน 3 หลัก "พ.ศ. 2570" จึงไม่โดน)
    text = re.sub("(?<![0-9๐-๙,.])(" + _DIG + "{1,3})" + _SEP + "(" + "|".join(UNITS) + ")",
                  r"\1" + NBSP + r"\2", text)
    text = re.sub("(" + _B + "(?:จำนวน|รวม(?:" + ZWSP + "?เป็น)?))" + _SEP + "(?=" + _DIG + "{1,3}" + NBSP + ")",
                  lambda m: m.group(1).replace(ZWSP, "") + NBSP, text)
    # จำนวนเงินกับ "บาท" ต้องอยู่บรรทัดเดียวกัน
    text = re.sub("(" + _DIG + "[0-9๐-๙,.]*)" + _SEP + "บาท", r"\1" + NBSP + "บาท", text)
    # ห้ามตัดหลังวงเล็บเปิดหรือก่อนวงเล็บปิด
    text = re.sub(r"\(" + _Z, "(", text)
    text = re.sub(_Z + r"\)", ")", text)
    # เลขแบบ/เลขหนังสือ เช่น "302/2" ห้ามขาดตรง /
    text = re.sub("(" + _DIG + ")" + _Z + "/", r"\1/", text)
    # เลขที่หนังสือทั้งก้อน "ที่ นร ๐๗๐๒/ว ๒๓๘"
    text = re.sub("(?:" + _B + "ที่" + _SEP + ")?นร" + _SEP + _DIG + "[0-9๐-๙." + ZWSP + "]*/" + ZWSP + "*"
                  "(?:ว[ " + ZWSP + "]*" + _DIG + "+(?:" + ZWSP + "*/" + ZWSP + "*" + _DIG + "+)?)?",
                  lambda m: m.group(0).replace(ZWSP, "").replace(" ", NBSP), text)
    # "ความ" ติดคำถัดไปเสมอ · "การ" แบบเดียวกันแต่ต้องยืนเป็นคำเดี่ยว ("ดำเนินการ" ไม่ติด)
    text = re.sub("ความ" + _Z, "ความ", text)
    text = re.sub(_B + "การ" + _Z, "การ", text)
    # วงเล็บสั้น ๆ (ไม่เกิน 12 ตัว) อยู่บรรทัดเดียวทั้งก้อน — "(เฟส 2)"
    text = re.sub(r"\(([^()]{1,12})\)",
                  lambda m: "(" + m.group(1).replace(ZWSP, "").replace(" ", NBSP) + ")", text)
    text = _glue_lead_words(text)
    text = text.replace("พ.ศ. ", "พ.ศ." + NBSP)
    # NBSP ที่มี ZWSP ตามหลังก็ยังตัดบรรทัดได้ — ลบ ZWSP ทิ้ง
    text = re.sub(NBSP + _Z, NBSP, text)
    return _reading_rules(text)


# ═══════════════════════════════════════════════════════════ ใช้งาน

def insert_zwsp(text: str, *, min_token: int = MIN_TOKEN, max_chunk: int = MAX_CHUNK) -> str:
    """ใส่ ZWSP + keep_together ให้ข้อความเดี่ยว ๆ (ถ้ามี run ให้ใช้ insert_zwsp_paragraph)"""
    return keep_together(_insert_raw(text, min_token=min_token, max_chunk=max_chunk))


def plain_text(text: str) -> str:
    """ข้อความก่อนผ่านตัวตัดคำ — ลบ ZWSP และคืน NBSP เป็นช่องว่าง (ไว้เทียบว่าข้อความไม่เปลี่ยน)"""
    return text.replace(ZWSP, "").replace(NBSP, " ")


def strip_zwsp(text: str) -> str:
    """เอา ZWSP ออกอย่างเดียว (NBSP ที่มัดคำไว้ยังอยู่)"""
    return text.replace(ZWSP, "")


def _map_back(runs, texts: list[str], out: str) -> int:
    """
    แมปข้อความที่ผ่านตัวตัดคำกลับเข้าแต่ละ run
    out ต่างจากต้นฉบับได้แค่ มี/ไม่มี ZWSP และช่องว่างที่กลายเป็น NBSP
    """
    full = "".join(texts)
    owner = [i for i, t in enumerate(texts) for _ in t]
    buf = [[] for _ in runs]
    j, pending, added = 0, [], 0
    for ch in out:
        if ch == ZWSP:
            pending.append(ch)
            continue
        if j >= len(full):
            raise ValueError("ข้อความยาวกว่าต้นฉบับ")
        if ch != full[j] and not (ch == NBSP and full[j] == " "):
            if ch == NBSP:
                continue                    # NBSP ที่ต้นฉบับไม่มีช่องว่าง — ทิ้ง ข้อความมาก่อน
            raise ValueError("ข้อความเปลี่ยนที่ตำแหน่ง %d: %r -> %r" % (j, full[j], ch))
        k = owner[j]
        buf[k].extend(pending)
        added += len(pending)
        pending = []
        buf[k].append(ch)
        j += 1
    if j != len(full):
        raise ValueError("แมปข้อความไม่ครบ %d/%d" % (j, len(full)))
    for r, b in zip(runs, buf):
        r.text = "".join(b)
    return added


def insert_zwsp_paragraph(paragraph, *, min_token: int = MIN_TOKEN,
                          max_chunk: int = MAX_CHUNK) -> int:
    """
    ใส่ ZWSP ให้ย่อหน้าของ python-docx โดยตัดคำจากข้อความรวมทั้งย่อหน้า
    แล้วแมปกลับเข้าแต่ละ run — กันรอยต่อ run ที่ผ่ากลางคำ

    คืนจำนวน ZWSP ที่ใส่
    """
    runs = [r for r in paragraph.runs if r.text]
    if not runs:
        return 0
    texts = [r.text.replace(ZWSP, "") for r in runs]
    out = insert_zwsp("".join(texts), min_token=min_token, max_chunk=max_chunk)
    return _map_back(runs, texts, out)


_JUSTIFIED = {"both", "distribute", "thaiDistribute", "lowKashida", "mediumKashida", "highKashida"}


def _alignment(paragraph):
    """ค่าการจัดวางที่มีผลจริง — ถ้าย่อหน้าไม่ได้ตั้งเอง ไล่หาจาก style"""
    from docx.oxml.ns import qn
    el = paragraph._element
    ppr = el.pPr
    jc = ppr.find(qn("w:jc")) if ppr is not None else None
    if jc is not None:
        return jc.get(qn("w:val"))
    style = paragraph.style
    while style is not None:
        sppr = style.element.pPr
        sjc = sppr.find(qn("w:jc")) if sppr is not None else None
        if sjc is not None:
            return sjc.get(qn("w:val"))
        style = style.base_style
    return None


def is_justified(paragraph) -> bool:
    return _alignment(paragraph) in _JUSTIFIED


def insert_zwsp_document(doc, *, only_justified: bool = True, include_tables: bool = False,
                         **kw) -> int:
    """
    ใส่ ZWSP ทั้งเอกสาร — ค่าตั้งต้นใส่เฉพาะย่อหน้าที่จัดชิดขอบขวา ตามที่โม 09 ทำ
    หัวเรื่อง ป้าย ลายเซ็น ตาราง ไม่ต้องมี ZWSP (บทเรียน ก.ค. 2569: ใส่ทุกย่อหน้าแล้ว "แยกถูกผิดไม่ออก")

    ถ้าไม่มีย่อหน้าไหนจัดชิดขอบขวาเลย จะเตือน แทนที่จะไม่ทำอะไรเงียบ ๆ
    """
    def targets():
        for p in doc.paragraphs:
            yield p
        if include_tables:
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs

    n = done = 0
    for p in targets():
        if only_justified and not is_justified(p):
            continue
        done += 1
        n += insert_zwsp_paragraph(p, **kw)
    if only_justified and done == 0:
        warnings.warn(
            "ไม่มีย่อหน้าที่จัดชิดขอบขวา จึงไม่ได้ใส่ ZWSP เลย — จัดชิดขอบขวาก่อน "
            "(thai_fit.justify_body) หรือส่ง only_justified=False | no justified paragraphs: "
            "nothing was changed", RuntimeWarning, stacklevel=2)
    return n


# ═══════════════════════════════════════════════════════════ ตรวจผลจริง

def find_midword_breaks(line_pairs) -> list[tuple[str, str]]:
    """
    รับคู่ (ท้ายบรรทัดบน, ต้นบรรทัดล่าง) ที่อ่านจาก PDF จริง คืนคู่ที่ตัดกลางคำ

    ข้อจำกัด: ใช้พจนานุกรมเดียวกับตัวตัดคำ จึงมองไม่เห็นคำที่พจนานุกรมไม่มี
    (เฝ้า|ระวัง · กรมควบคุม|โรค) — ต้องอ่านรอยต่อบรรทัดทุกจุดเองด้วย inspect_docx --breaks
    """
    bad = []
    for tail, head in line_pairs:
        tail, head = tail.rstrip(), head.lstrip()
        if not tail or not head:
            continue
        if not (is_thai(tail[-1]) and is_thai(head[0])):
            continue
        joined = tail + head
        pos, cuts = 0, set()
        for tok in _tokenize(joined):
            pos += len(tok)
            cuts.add(pos)
        if len(tail) not in cuts:
            bad.append((tail[-14:], head[:14]))
    return bad
