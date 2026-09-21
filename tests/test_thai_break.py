# -*- coding: utf-8 -*-
"""
เทสต์ตัวตัดคำ — ทุกข้อมาจากจุดที่เคยพังจริงในเอกสารที่ส่งไปแล้ว
รัน: python -m pytest tests -q   หรือ   python tests/test_thai_break.py
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", "scripts"))
import thai_break as tb  # noqa: E402

Z = tb.ZWSP


def cut(text):
    """ใส่ ZWSP แล้วแทนด้วย | เพื่ออ่านง่ายในเทสต์"""
    return tb.insert_zwsp(text).replace(Z, "|")


# ─────────────────────────────────── ข้อความจริงจากเอกสารที่เคยส่ง

VACCINE = ("สถาบันวัคซีนแห่งชาติได้รับงบประมาณรายจ่ายประจำปีงบประมาณ พ.ศ. 2570 "
           "เพื่อดำเนินการตามแผนงานโครงการ")
KPI = ("ขอเรียนชี้แจงว่าการดำเนินการตามตัวชี้วัดและการก่อหนี้ผูกพัน"
       "เป็นไปตามแนวทางที่สำนักงบประมาณกำหนด")
MONEY = ("มีวงเงินงบประมาณรวมทั้งสิ้น 50,038,500 บาท ตามที่ขอทำความตกลงไปได้ "
         "ความละเอียดแจ้งแล้ว")


def test_text_unchanged():
    """ZWSP ต้องไม่เปลี่ยนข้อความ — ถอดออกแล้วต้องได้ของเดิมเป๊ะ"""
    for t in (VACCINE, KPI, MONEY):
        assert tb.strip_zwsp(tb.insert_zwsp(t)) == t


def test_tail_word_never_starts_line():
    """'แห่งชาติ' ห้ามขึ้นต้นบรรทัด — เคยหลุดเพราะการ glue กระโดดข้ามการตรวจ"""
    assert "วัคซีน|แห่งชาติ" not in cut(VACCINE)


def test_keep_together_phrases():
    """วลีที่กิ๊ฟสั่งมัดไว้ ห้ามมีจุดตัดคั่นกลาง"""
    assert "ประจำปีงบประ|มาณ" not in cut(VACCINE)     # เคยถูกตัวซอยผ่า
    assert "ตัว|ชี้" not in cut(KPI) and "ชี้|วัด" not in cut(KPI)
    assert "ก่อ|หนี้" not in cut(KPI)
    assert "แนว|ทาง" not in cut(KPI)
    assert "รวม|ทั้งสิ้น" not in cut(MONEY)
    assert "ความ|ละเอียด" not in cut(MONEY)


def test_nominal_prefix_stays_with_verb():
    """ตัวตัดคำชอบผลัก 'การ' ไปเกาะคำถัดไป — 'ดำเนิน|การตาม' ต้องกลายเป็น 'ดำเนินการ|ตาม'"""
    assert "ดำเนิน|การ" not in cut(VACCINE)


def test_numbers_and_latin_untouched():
    """ใส่เฉพาะจุดไทยชนไทย ตัวเลขและอังกฤษต้องไม่ถูกแหก"""
    assert "50,038,500" in cut(MONEY)
    t = "ระบบ VIMS บนคลาวด์ GDCC รองรับ Mobile device ของหน่วยงาน"
    assert "VIMS" in cut(t) and "Mobile" in cut(t) and "GDCC" in cut(t)


def test_phone_number_intact():
    """เบอร์โทรห้ามขาด"""
    assert "012-345-6789" in cut("ผู้ชี้แจงหมายเลขโทรศัพท์ 012-345-6789 ติดต่อได้")


def test_split_long_chunks_uses_dictionary():
    """ซอยคำยาวต้องได้คำจริงทั้งสองฝั่ง ไม่ใช่ซอยตรงกลางมั่ว ๆ"""
    assert tb.split_long_chunks(["งบประมาณรายจ่าย"]) == ["งบประมาณ", "รายจ่าย"]
    # วลีที่มัดไว้ห้ามโดนซอย แม้จะยาวเกิน MAX_CHUNK
    assert tb.split_long_chunks(["ประจำปีงบประมาณ"]) == ["ประจำปีงบประมาณ"]


def test_no_break_before_vowel_marks():
    """ห้ามขึ้นบรรทัดใหม่ด้วยสระบน/ล่าง/วรรณยุกต์"""
    out = tb.insert_zwsp(KPI)
    for i, ch in enumerate(out):
        if ch == Z and i + 1 < len(out):
            assert out[i + 1] not in tb.NO_BREAK_BEFORE


def test_leading_vowels_are_breakable():
    """
    สระหน้า เ แ โ ใ ไ ต้องตัดข้างหน้าได้ — ถ้าห้าม คำอย่าง และ/ได้/เป็น
    จะเกาะคำก่อนหน้าเป็นก้อนยาว 25-30 ตัว แล้วบรรทัดจะถ่าง
    """
    for ch in "เแโใไ":
        assert ch not in tb.NO_BREAK_BEFORE
    assert "|และ" in cut("การดำเนินงานและการติดตามผลการปฏิบัติงานของหน่วยรับงบประมาณ")


def test_paragraph_level_beats_per_run():
    """
    หัวใจของรีโปนี้: เอกสารจาก Word มี run ที่ผ่ากลางคำ
    ตัดคำทีละ run จะได้ขอบเขตผิด ('โด' + 'น' -> ตัดกลาง 'โดน')
    """
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("ประชาชนไม่โด")      # run ผ่ากลางคำแบบที่เจอจริง
    p.add_run("นทอดทิ้ง")
    tb.insert_zwsp_paragraph(p)
    joined = "".join(r.text for r in p.runs).replace(Z, "|")
    assert "โด|น" not in joined
    assert tb.strip_zwsp("".join(r.text for r in p.runs)) == "ประชาชนไม่โดนทอดทิ้ง"


if __name__ == "__main__":
    import io
    # คอนโซล Windows เป็น cp874/cp1252 ภาษาไทยจะเพี้ยน ต้องบังคับ utf-8
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print("  ผ่าน  %s" % name)
            except AssertionError as e:
                fails += 1
                print("  ตก    %s  %s" % (name, e))
    print("\n%s" % ("ผ่านทั้งหมด" if not fails else "ตก %d ข้อ" % fails))
    sys.exit(1 if fails else 0)
