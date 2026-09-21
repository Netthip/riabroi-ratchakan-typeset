# -*- coding: utf-8 -*-
"""
เทสต์ส่วนที่ขึ้นกับเครื่อง (Windows / macOS / LibreOffice) — จำลอง osascript และ soffice ด้วย mock
จึงรันได้ทุกเครื่อง
รัน: python -m pytest tests -q   หรือ   python tests/test_platform.py

ทางเดิน Word for Mac ทดสอบกับ Word จริงบนเครื่อง Windows ไม่ได้ — เทสต์ชุดนี้ล็อกกติกาความปลอดภัยไว้
(เปิดแต่สำเนา ไม่แตะเอกสารที่ผู้ใช้เปิดค้าง ไม่ปิด Word ของผู้ใช้) ไม่ให้ใครแก้หลุดโดยไม่รู้ตัว
"""
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import inspect_docx as ix  # noqa: E402

MAC_WORD = "/Applications/Microsoft Word.app"


def _docx(path, font="TH SarabunPSK"):
    from docx import Document
    import thai_fit
    doc = Document()
    thai_fit.set_run_font(doc.add_paragraph().add_run("ทดสอบฟอนต์"), font, 16)
    doc.save(path)


def _pdf(path, *, creator="", producer=""):
    import fitz
    d = fitz.open()
    d.new_page().insert_text((72, 72), "test", fontname="helv")
    d.set_metadata({"creator": creator, "producer": producer})
    d.save(path)
    d.close()


def _mac(work):
    """จำลองเครื่อง Mac ที่มี Word"""
    return [mock.patch.object(sys, "platform", "darwin"),
            mock.patch.dict(os.environ, {"RIABROI_WORKDIR": work}),
            mock.patch.object(ix, "find_word", return_value=MAC_WORD)]


def _run_with(patches, fn):
    for p in patches:
        p.start()
    try:
        return fn()
    finally:
        for p in reversed(patches):
            p.stop()


def _leftovers(work):
    return [f for f in os.listdir(work) if f.startswith(("riabroi-", "~$"))]


# ─────────────────────────────────────────────────── Word for Mac (จำลอง)

def test_mac_opens_private_copy_and_cleans_up():
    """เปิดแต่สำเนาชื่อสุ่มในโฟลเดอร์ทำงาน · ได้ PDF ที่ปลายทาง · ไม่มีไฟล์ค้าง"""
    work, out = tempfile.mkdtemp(), tempfile.mkdtemp()
    src = os.path.join(out, "ต้นฉบับ.docx")
    _docx(src)

    def fake_osascript(cmd, **kw):
        assert cmd[:2] == ["osascript", ix.MAC_SCRIPT]
        copy_docx, copy_pdf, stem = cmd[2], cmd[3], cmd[4]
        assert os.path.dirname(copy_docx) == os.path.abspath(work)     # ไม่ใช่ไฟล์จริงของผู้ใช้
        assert os.path.basename(copy_docx) == stem + ".docx" and stem in copy_pdf
        with open(copy_docx, "rb") as a, open(src, "rb") as b:
            assert a.read() == b.read()
        _pdf(copy_pdf)
        open(os.path.join(work, "~$" + os.path.basename(copy_docx)[2:]), "w").close()
        return subprocess.CompletedProcess(cmd, 0, "", "")

    try:
        pdf = _run_with(_mac(work) + [mock.patch.object(ix.subprocess, "run", side_effect=fake_osascript)],
                        lambda: ix.to_pdf(src, out))
        assert pdf == os.path.join(out, "ต้นฉบับ.pdf") and os.path.isfile(pdf)
        assert _leftovers(work) == [], _leftovers(work)
    finally:
        shutil.rmtree(work, ignore_errors=True)
        shutil.rmtree(out, ignore_errors=True)


def test_mac_failure_explains_and_cleans_up():
    """macOS ไม่ให้สิทธิ์ (-1743) ต้องบอกวิธีแก้เป็นภาษาไทย และไม่ทิ้งสำเนาไว้"""
    work, out = tempfile.mkdtemp(), tempfile.mkdtemp()
    src = os.path.join(out, "a.docx")
    _docx(src)
    denied = subprocess.CompletedProcess(
        [], 1, "", "execution error: Not authorized to send Apple events to Microsoft Word. (-1743)")
    try:
        try:
            _run_with(_mac(work) + [mock.patch.object(ix.subprocess, "run", return_value=denied)],
                      lambda: ix.to_pdf(src, out))
            raise AssertionError("ต้องล้มเมื่อ osascript ล้ม")
        except RuntimeError as e:
            assert "Automation" in str(e) and "-1743" in str(e)
        assert _leftovers(work) == []
        assert not os.path.exists(os.path.join(out, "a.pdf"))
    finally:
        shutil.rmtree(work, ignore_errors=True)
        shutil.rmtree(out, ignore_errors=True)


def test_mac_error_hints():
    assert "Grant File Access" in ix._mac_error("execution error: AppleEvent timed out. (-1712)")
    unknown = ix._mac_error("execution error: something new (-42)")
    assert "ข้อความจาก osascript" in unknown and "-42" in unknown


def test_applescript_never_touches_other_documents():
    """กติกาความปลอดภัยใน AppleScript — แก้สคริปต์แล้วข้อไหนหลุด เทสต์นี้ต้องตก"""
    with open(ix.MAC_SCRIPT, encoding="utf-8") as f:
        src = f.read()
    assert all(ord(c) < 128 for c in src), "ต้องเป็น ASCII — osascript อาจไม่อ่านไฟล์เป็น UTF-8"
    code = [ln.split("--")[0] for ln in src.splitlines()]           # ตัดคอมเมนต์
    body = "\n".join(code)
    assert "on run argv" in body and "format PDF" in body
    assert "active document" not in body                             # อ้างด้วยชื่อสำเนาเท่านั้น
    for bad in ("do shell script", "killall", "every window"):
        assert bad not in body, bad
    closes = [ln for ln in code if "close " in ln]
    assert closes and all("docStem" in ln and "saving no" in ln for ln in closes)
    quits = [ln for ln in code if "quit" in ln]
    assert quits and all("wasRunning" in ln and "count of documents" in ln for ln in quits)


# ─────────────────────────────────────────────────── เลือกตัวส่งออก

def test_auto_without_word_uses_libreoffice_and_warns():
    out = tempfile.mkdtemp()
    src = os.path.join(out, "a.docx")
    _docx(src)

    def fake_soffice(cmd, **kw):
        assert cmd[0] == "soffice" and "--headless" in cmd
        assert any(c.startswith("-env:UserInstallation=file:") for c in cmd)   # โปรไฟล์แยก
        _pdf(os.path.join(cmd[cmd.index("--outdir") + 1], "a.pdf"),
             creator="Writer", producer="LibreOffice 24.2")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    try:
        with mock.patch.object(ix, "find_word", return_value=None), \
                mock.patch.object(ix, "find_libreoffice", return_value="soffice"), \
                mock.patch.object(ix.subprocess, "run", side_effect=fake_soffice), \
                warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            pdf = ix.to_pdf(src, out)
        assert os.path.isfile(pdf) and ix.pdf_engine(pdf) == "libreoffice"
        assert any("LibreOffice" in str(w.message) for w in caught)
    finally:
        shutil.rmtree(out, ignore_errors=True)


def test_word_engine_without_word_is_an_error():
    out = tempfile.mkdtemp()
    try:
        with mock.patch.object(ix, "find_word", return_value=None):
            try:
                ix.to_pdf(os.path.join(out, "a.docx"), out, engine="word")
                raise AssertionError("ไม่มี Word ต้องล้ม ไม่ใช่แอบไปใช้ LibreOffice")
            except RuntimeError as e:
                assert "Word" in str(e)
    finally:
        shutil.rmtree(out, ignore_errors=True)


def test_engine_read_from_pdf_metadata():
    tmp = tempfile.mkdtemp()
    cases = {"word": ("Microsoft® Word for Microsoft 365", "Microsoft® Word for Microsoft 365"),
             "word-mac": ("Microsoft Word", "macOS Version 14.5 (Build 23F79) Quartz PDFContext"),
             "libreoffice": ("Writer", "LibreOffice 24.2")}
    try:
        for want, (creator, producer) in cases.items():
            p = os.path.join(tmp, want + ".pdf")
            _pdf(p, creator=creator, producer=producer)
            assert ix.pdf_engine(p) == want, (want, ix.pdf_engine(p))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ─────────────────────────────────────────────────── ฟอนต์ถูกแทน

def test_font_names_compare_across_docx_and_pdf():
    n = ix._norm_font
    assert n("BCDEEE+THSarabunPSK-Bold") == n("TH SarabunPSK")
    assert n("THSarabunNew-BoldItalic") == n("TH Sarabun New")
    assert n("THSarabunIT9") == n("TH SarabunIT" + chr(0x0E59))           # IT๙
    assert n("TimesNewRomanPS-BoldMT") == n("Times New Roman")
    assert n("ArialMT") == n("Arial")
    assert n("Thonburi") != n("TH SarabunPSK")


def test_font_check_flags_missing_thai_font():
    """เอกสารสั่ง TH SarabunPSK แต่ PDF มีแค่ Helvetica = เครื่องไม่มีฟอนต์ ผลวัดใช้ไม่ได้"""
    tmp = tempfile.mkdtemp()
    try:
        src, pdf = os.path.join(tmp, "a.docx"), os.path.join(tmp, "a.pdf")
        _docx(src)
        _pdf(pdf)
        fc = ix.font_check(src, pdf)
        assert fc["main"] == "TH SarabunPSK" and fc["main_missing"]
        assert "Helvetica" in fc["others"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_lang_code_is_not_a_font():
    """w:eastAsia ใน <w:lang> เป็นรหัสภาษา ห้ามนับเป็นฟอนต์"""
    xml = ('<w:rPr><w:rFonts w:ascii="TH SarabunPSK" w:cs="TH SarabunPSK"/>'
           '<w:lang w:val="en-US" w:eastAsia="en-US" w:bidi="th-TH"/></w:rPr>')
    assert ix._rfonts(xml) == ["TH SarabunPSK", "TH SarabunPSK"]


def _font_with_names(records):
    """ไฟล์ฟอนต์จิ๋วที่มีแค่ตาราง name — พอให้ทดสอบตัวอ่านชื่อโดยไม่ต้องพกไฟล์ฟอนต์จริง"""
    import struct
    strings, recs = b"", b""
    for pid, lang, nid, text in records:
        raw = text.encode("mac_roman" if pid == 1 else "utf-16-be")
        recs += struct.pack(">HHHHHH", pid, 0 if pid == 1 else 1, lang, nid, len(raw), len(strings))
        strings += raw
    table = struct.pack(">HHH", 0, len(records), 6 + len(recs)) + recs + strings
    return (struct.pack(">IHHHH", 0x00010000, 1, 16, 0, 0)
            + struct.pack(">4sIII", b"name", 0, 28, len(table)) + table)


def test_font_family_it9_claims_psk_on_mac():
    """ไฟล์ TH SarabunIT๙ รุ่น Windows: Windows เห็น IT๙ แต่ชื่อที่ Mac ใช้จัดกลุ่มคือ PSK (วัดจากไฟล์จริง)"""
    import doctor
    it9 = "TH SarabunIT" + chr(0x0E59)
    data = _font_with_names([(1, 0, 1, "TH SarabunPSK"), (1, 0, 16, "TH SarabunPSK"),
                             (3, 0x409, 1, it9), (3, 0x409, 16, "TH SarabunPSK")])
    assert doctor.font_family(data) == it9
    assert doctor.font_family(data, mac=True) == "TH SarabunPSK"
    psk = _font_with_names([(3, 0x409, 1, "TH SarabunPSK")])
    assert doctor.font_family(psk) == doctor.font_family(psk, mac=True) == "TH SarabunPSK"
    assert doctor.font_family(b"not a font") is None


def test_breaks_where_allowed():
    """ตัวตรวจเครื่อง: Word ตัดตรง ZWSP หรือช่องว่างได้ · ตรง NBSP หรือกลางคำไม่ได้"""
    import doctor
    z, nb = chr(0x200B), chr(0x00A0)
    src = "สถาบัน" + z + "วัคซีน" + z + "แห่งชาติ ได้รับ" + nb + "งบประมาณ"
    cases = [(["สถาบันวัคซีน", "แห่งชาติ ได้รับ งบประมาณ"], (1, 1)),     # ตรง ZWSP
             (["สถาบันวัคซีนแห่งชาติ", "ได้รับ งบประมาณ"], (1, 1)),       # ตรงช่องว่าง
             (["สถาบันวัค", "ซีนแห่งชาติ ได้รับ งบประมาณ"], (0, 1)),       # กลางคำ
             (["สถาบันวัคซีนแห่งชาติ ได้รับ", "งบประมาณ"], (0, 1)),       # ตรง NBSP
             (["ข้อความอื่น"], None)]                                      # จับคู่ไม่ติด
    for lines, want in cases:
        with mock.patch.object(doctor, "_visual_lines", return_value=lines):
            assert doctor.breaks_where_allowed(src, "x.pdf") == want, (lines, want)


# ─────────────────────────────────────────────────── อักษรเพี้ยนจาก PDF

def test_pdf_text_repair():
    tho_low = chr(0xF70B)                         # ไม้โทตัวต่ำ (PUA ของ TH SarabunPSK)
    tho, nikhahit, aa, am = chr(0x0E49), chr(0x0E4D), chr(0x0E32), chr(0x0E33)
    assert ix.fix_thai_text("น" + tho_low + nikhahit + aa) == "น" + tho + am      # น้ำ
    assert ix.fix_thai_text("ด" + " " + aa + "เนิน") == "ด" + am + "เนิน"         # ด าเนิน
    assert ix.fix_thai_text("ประจำปีงบประมาณ พ.ศ. 2570") == "ประจำปีงบประมาณ พ.ศ. 2570"


# ─────────────────────────────────────────────────── ติดตั้ง

def test_pywin32_only_on_windows():
    """pywin32 ไม่มีบน Mac — ไม่ใส่เงื่อนไข pip install ล้มทั้งไฟล์"""
    with open(os.path.join(ROOT, "requirements.txt"), encoding="utf-8") as f:
        lines = [ln.strip() for ln in f if ln.lower().startswith("pywin32")]
    assert lines and all('sys_platform == "win32"' in ln for ln in lines)


if __name__ == "__main__":
    import io
    # คอนโซล Windows เป็น cp874/cp1252 ภาษาไทยจะเพี้ยน ต้องบังคับ utf-8
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print("  ผ่าน  %s" % name)
            except Exception as e:  # noqa: BLE001 — รายงานทุกแบบ ไม่ให้ข้อเดียวล้มทั้งชุด
                fails += 1
                print("  ตก    %s  %s: %s" % (name, type(e).__name__, e))
    print("\n%s" % ("ผ่านทั้งหมด" if not fails else "ตก %d ข้อ" % fails))
    sys.exit(1 if fails else 0)
